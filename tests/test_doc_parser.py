"""Document upload parser: PDF, Word, text."""

import io

import docx
import pymupdf
import pytest

from app.bot import kb
from app.services import doc_parser
from app.services.doc_parser import ParseError

BODY = (
    "Pengobatan TBC sensitif obat berlangsung minimal enam bulan dan obatnya tersedia "
    "gratis di puskesmas terdekat. Minum obat setiap hari pada jam yang sama."
)


def make_pdf(pages: int = 3, scanned: bool = False) -> bytes:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        if scanned:  # an image only, no text layer
            page.insert_image(
                pymupdf.Rect(50, 50, 300, 300),
                pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), 0),
            )
            continue
        y = 80
        if i == 0:
            page.insert_text((60, y), "Panduan Pasien TBC", fontsize=22, fontname="hebo")
            y += 50
        page.insert_text((60, y), f"{i + 1}. Bagian ke-{i + 1}", fontsize=14, fontname="hebo")
        y += 30
        tb = pymupdf.Rect(60, y, 535, y + 120)
        page.insert_textbox(
            tb, BODY + " Kepatuhan minum obat sangat penting bagi pengo-\nbatan.", fontsize=11
        )
        y += 140
        page.insert_textbox(
            pymupdf.Rect(60, y, 535, y + 60),
            "- Bawa kartu berobat\n- Datang tepat waktu",  # "•" isn't in the test font
            fontsize=11,
        )
        page.insert_text((60, 815), "Yayasan Contoh TBC Indonesia", fontsize=8)  # running footer
        page.insert_text((520, 815), str(i + 1), fontsize=8)  # page number
    return doc.tobytes()


def make_docx() -> bytes:
    d = docx.Document()
    d.core_properties.title = "SOP Pendampingan Pasien"
    d.add_paragraph("Dokumen ini menjelaskan langkah pendampingan pasien TBC oleh kader.")
    d.add_heading("Kunjungan rumah", level=1)
    d.add_paragraph(BODY)
    d.add_paragraph("Bawa formulir kunjungan", style="List Bullet")
    d.add_paragraph("Catat keluhan pasien", style="List Bullet")
    d.add_heading("Jadwal kontrol", level=2)
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate(
        [("Bulan", "Kegiatan"), ("1", "Kontrol mingguan"), ("2-6", "Kontrol bulanan")]
    ):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_structure_and_cleanup():
    p = doc_parser.parse(make_pdf(), "panduan.pdf")
    assert p.title == "Panduan Pasien TBC" and p.pages == 3
    assert "## 1. Bagian ke-1" in p.markdown and p.sections == 3
    assert "pengobatan." in p.markdown  # hyphenation rejoined
    assert "- Bawa kartu berobat\n- Datang tepat waktu" in p.markdown
    assert "Yayasan Contoh" not in p.markdown  # running footer dropped
    assert "\n\n3\n\n" not in p.markdown and not p.markdown.rstrip().endswith("3")


def test_scanned_pdf_is_refused():
    with pytest.raises(ParseError, match="scan"):
        doc_parser.parse(make_pdf(scanned=True), "scan.pdf")


def test_docx_headings_lists_tables():
    p = doc_parser.parse(make_docx(), "sop.docx")
    assert p.title == "SOP Pendampingan Pasien"
    assert p.markdown.startswith("Dokumen ini menjelaskan")  # intro before the first heading
    assert "## Kunjungan rumah" in p.markdown and "### Jadwal kontrol" in p.markdown
    assert "- Bawa formulir kunjungan\n- Catat keluhan pasien" in p.markdown
    assert "- Bulan: 1; Kegiatan: Kontrol mingguan" in p.markdown


def test_markdown_and_text():
    md = "# Judul\n\nPembuka singkat tentang TBC dan pengobatannya.\n\n## Gejala\n\nBatuk dua minggu atau lebih."
    p = doc_parser.parse(md.encode(), "catatan.md")
    assert p.title == "Judul" and p.markdown.startswith("Pembuka") and "## Gejala" in p.markdown


def test_unsupported_and_empty():
    with pytest.raises(ParseError, match="PDF, Word"):
        doc_parser.parse(b"x", "foto.jpg")
    with pytest.raises(ParseError):
        doc_parser.parse(b"   ", "kosong.txt")


def test_output_is_indexable(tmp_path):
    """The bot's indexer must accept what the parser produces."""
    p = doc_parser.parse(make_docx(), "sop.docx")
    f = tmp_path / "x.md"
    f.write_text(
        f"---\nid: p99\ntitle: {p.title}\nshort_cite: {p.title}\n---\n{p.markdown}\n",
        encoding="utf-8",
    )
    _paper, parsed = kb.parse_markdown(f)
    headings = [s.heading for s in parsed.sections]
    assert "Kunjungan rumah" in headings and "Jadwal kontrol" in headings
